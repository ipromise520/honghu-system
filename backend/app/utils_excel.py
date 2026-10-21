"""Excel / CSV 上传解析工具"""
import csv
import io


async def parse_upload_sheets(request, sheet_names=None, skip_rows=0) -> tuple[str, dict[str, list[dict]]]:
    """解析 multipart 请求中的 xlsx / csv。

    sheet_names 为白名单（手册 10.2：只读取需要的 sheet，如「当月明细」），
    None 表示读取全部 sheet。skip_rows 为每个 sheet 表头前的标题行数
    （如 H5 通报表首行为标题、次行才是列名则传 1）；也可传 dict {sheet名: 跳行数}
    按 sheet 分别指定（未列出的 sheet 用默认跳行数，dict 模式下默认为 1），
    用于同一文件内表头行不一致的场景（如 H5 新格式：「装机回单」表头在第 1 行、其余在第 2 行）。
    返回 (文件名, {sheet名: 行列表[dict]})；csv 返回 {"": 行列表}。
    """
    import openpyxl

    ctype = request.headers.get("content-type", "")
    if "multipart/form-data" in ctype:
        form = await request.form()
        upload = form.get("file")
        if upload is None:
            raise ValueError("缺少 file 字段")
        filename = upload.filename or "upload"
        data = await upload.read()
    else:
        filename = "raw.csv"
        data = await request.body()

    # skip_rows 归一化：dict 模式（按 sheet 指定，默认 1）或 int 模式（全体同一跳行数）
    skip_map = skip_rows if isinstance(skip_rows, dict) else {}
    default_skip = skip_rows if isinstance(skip_rows, int) else 1

    # .xls 后缀但实际为 xlsx 格式（导出系统常见），统一用 openpyxl 尝试
    if filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        # read_only=True：流式读取，避免报表 sheet 声明整表尺寸（如 1048576 行）时
        # 非流式模式构建海量空单元格导致解析极慢/超时（大体积通报表常见）
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        wanted = {str(s).strip() for s in sheet_names} if sheet_names else None
        out: dict[str, list[dict]] = {}
        for ws in wb.worksheets:
            if wanted is not None and ws.title.strip() not in wanted:
                continue
            n_skip = max(0, int(skip_map.get(ws.title.strip(), default_skip)))
            it = ws.iter_rows(values_only=True)
            for _ in range(n_skip):                  # 跳过表头前的标题行
                next(it, None)
            header = next(it, None)
            if not header:
                continue
            # 重复列名加 .N 后缀（与 pandas 一致，避免同行同名列相互覆盖丢值）
            seen: dict = {}
            norm_header = []
            for h in header:
                h = str(h).strip() if h is not None else ""
                if h in seen:
                    seen[h] += 1
                    norm_header.append(f"{h}.{seen[h]}")
                else:
                    seen[h] = 0
                    norm_header.append(h)
            header = norm_header
            rows = []
            try:
                for r in it:
                    if all(v is None or str(v).strip() == "" for v in r):
                        continue
                    rows.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
            except ValueError:
                # 规避 openpyxl 解析表尾 autoFilter（自动筛选）的已知缺陷：
                # 筛选值非法时在数据行全部输出后抛 ValueError，数据不受影响
                pass
            out[ws.title] = rows
        wb.close()
        return filename, out

    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    all_rows = list(reader)
    if not all_rows:
        return filename, {"": []}
    header = [h.strip() for h in all_rows[0]]
    return filename, {"": [{header[i]: row[i] if i < len(row) else "" for i in range(len(header))}
                           for row in all_rows[1:] if any(str(c).strip() for c in row)]}


async def parse_upload_raw(request, sheet_names=None) -> tuple[str, dict[str, list[tuple]]]:
    """解析 multipart 请求中的 xlsx，返回 (文件名, {sheet名: 原始行元组列表})。

    与 parse_upload_sheets 不同：不设表头、不跳行，保留完整矩阵，
    供多行（合并）表头的通报表解析使用（如 AI装维即访 的「装后即访-通报表」）。
    """
    import openpyxl

    ctype = request.headers.get("content-type", "")
    if "multipart/form-data" in ctype:
        form = await request.form()
        upload = form.get("file")
        if upload is None:
            raise ValueError("缺少 file 字段")
        filename = upload.filename or "upload"
        data = await upload.read()
    else:
        filename = "raw.csv"
        data = await request.body()

    if filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        # 同 parse_upload_sheets：流式读取，规避大尺寸报表 sheet 的解析卡顿
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        wanted = {str(s).strip() for s in sheet_names} if sheet_names else None
        out: dict[str, list[tuple]] = {}
        for ws in wb.worksheets:
            if wanted is not None and ws.title.strip() not in wanted:
                continue
            out[ws.title] = [tuple(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return filename, out

    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return filename, {"": [tuple(r) for r in reader]}


async def parse_upload(request) -> tuple[str, list[dict]]:
    """从 multipart 请求解析 xlsx / csv（全部 sheet 合并），返回 (文件名, 行列表[dict])"""
    filename, sheets = await parse_upload_sheets(request)
    rows: list[dict] = []
    for rs in sheets.values():
        rows.extend(rs)
    return filename, rows
